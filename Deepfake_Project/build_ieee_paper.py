from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Deepfake_Misinformation_Detection_IEEE_6_Page_Paper.docx"
FIG = ROOT / "paper_architecture.png"
TRAINING_FLOW_FIG = ROOT / "paper_training_flow.png"
DECISION_FLOW_FIG = ROOT / "paper_decision_flow.png"
FONT = "Times New Roman"


def font(run, size=10, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return run


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trPr.append(node)


def add_columns(section, count=2, space=360):
    sectPr = section._sectPr
    cols = sectPr.xpath("./w:cols")
    cols = cols[0] if cols else OxmlElement("w:cols")
    if not cols.getparent():
        sectPr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space))


def configure_section(section, columns=1):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    add_columns(section, columns)


def add_body(doc, text, first_indent=True):
    p = doc.add_paragraph(style="Body")
    p.paragraph_format.first_line_indent = Inches(0.15) if first_indent else None
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(text))
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="IEEE Heading")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(text.upper()), 10)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph(style="IEEE Subheading")
    font(p.add_run(text), 10, italic=True)
    return p


def add_equation(doc, text, number):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(text), 10, italic=True)
    font(p.add_run(f"     ({number})"), 10)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_header(table.rows[0])
    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        shade(cell, "D9E2F3")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(header), 8, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, (value, width) in enumerate(zip(row, widths)):
            cells[i].width = Inches(width)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(str(value)), 8)
    table.rows[-1]._tr.get_or_add_trPr()
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(text), 8)


def make_figure():
    canvas = Image.new("RGB", (1224, 900), "white")
    draw = ImageDraw.Draw(canvas)
    try:
        label_font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        label_font = ImageFont.load_default()
    boxes = [
        (45, 80, 315, 185, "Image\nUpload", "#E9EEFF"),
        (455, 80, 770, 185, "224 x 224\nPreprocessing", "#E9EEFF"),
        (865, 80, 1180, 185, "EfficientNet\nClassifier", "#DDEFEA"),
        (45, 360, 315, 465, "News\nText", "#FFF1D9"),
        (455, 360, 770, 465, "TF-IDF\nFeatures", "#FFF1D9"),
        (865, 360, 1180, 465, "Logistic\nClassifier", "#DDEFEA"),
        (455, 650, 770, 755, "Flask REST API", "#E8E8F3"),
        (865, 650, 1180, 755, "Android Client\n+ Web UI", "#E8E8F3"),
    ]
    for x1, y1, x2, y2, label, color in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=color,
                               outline="#334155", width=3)
        bbox = draw.multiline_textbbox((0, 0), label, font=label_font,
                                       align="center", spacing=5)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.multiline_text(((x1+x2-tw)/2, (y1+y2-th)/2), label,
                            fill="#172033", font=label_font, align="center", spacing=5)
    arrows = [((315, 132), (455, 132)), ((770, 132), (865, 132)),
              ((315, 412), (455, 412)), ((770, 412), (865, 412)),
              ((1020, 185), (650, 650)), ((1020, 465), (680, 650)),
              ((770, 702), (865, 702))]
    for start, end in arrows:
        draw.line((start, end), fill="#475569", width=5)
        ex, ey = end
        draw.polygon([(ex, ey), (ex-18, ey-10), (ex-18, ey+10)], fill="#475569")
    canvas.save(FIG)


def make_flowcharts():
    try:
        label_font = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        label_font = ImageFont.load_default()

    def horizontal_flow(labels, output, colors):
        canvas = Image.new("RGB", (1240, 360), "white")
        draw = ImageDraw.Draw(canvas)
        box_w, box_h, gap, x0, y = 210, 115, 38, 20, 115
        for index, label in enumerate(labels):
            x = x0 + index * (box_w + gap)
            draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=18,
                                   fill=colors[index], outline="#334155", width=3)
            bbox = draw.multiline_textbbox((0, 0), label, font=label_font,
                                           align="center", spacing=4)
            tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            draw.multiline_text((x + (box_w - tw) / 2, y + (box_h - th) / 2),
                                label, fill="#172033", font=label_font,
                                align="center", spacing=4)
            if index < len(labels) - 1:
                start_x, end_x, mid_y = x + box_w, x + box_w + gap, y + box_h // 2
                draw.line((start_x, mid_y, end_x, mid_y), fill="#475569", width=5)
                draw.polygon([(end_x, mid_y), (end_x - 15, mid_y - 9),
                              (end_x - 15, mid_y + 9)], fill="#475569")
        canvas.save(output)

    horizontal_flow(
        ["Collect and\nlabel data", "Stratified\n70/15/15 split", "Augment and\npreprocess", "Train and\nfine-tune", "Calibrate and\ntest"],
        TRAINING_FLOW_FIG,
        ["#E9EEFF", "#E9EEFF", "#FFF1D9", "#DDEFEA", "#E8E8F3"],
    )
    horizontal_flow(
        ["Image or\nnews input", "Preprocess\ninput", "Model\nprobability", "Confidence\ncheck", "Label or\nmanual review"],
        DECISION_FLOW_FIG,
        ["#E9EEFF", "#FFF1D9", "#DDEFEA", "#E8E8F3", "#E9EEFF"],
    )


def add_project_picture(doc, path, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    picture = p.add_run().add_picture(str(path), width=Inches(3.25))
    picture._inline.docPr.set("descr", alt)
    add_caption(doc, caption)


doc = Document()
configure_section(doc.sections[0], 1)
styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(10)

body = styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
body.font.name = FONT
body.font.size = Pt(10)
body.paragraph_format.space_before = Pt(0)
body.paragraph_format.space_after = Pt(3)
body.paragraph_format.line_spacing = 1.0

h = styles.add_style("IEEE Heading", WD_STYLE_TYPE.PARAGRAPH)
h.font.name = FONT
h.font.size = Pt(10)
h.paragraph_format.space_before = Pt(7)
h.paragraph_format.space_after = Pt(3)
h.paragraph_format.keep_with_next = True

sh = styles.add_style("IEEE Subheading", WD_STYLE_TYPE.PARAGRAPH)
sh.font.name = FONT
sh.font.size = Pt(10)
sh.font.italic = True
sh.paragraph_format.space_before = Pt(4)
sh.paragraph_format.space_after = Pt(1)
sh.paragraph_format.keep_with_next = True

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
font(title.add_run("A Lightweight Multimodal Framework for Deepfake Image and Fake-News Detection"), 20)

author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
author.paragraph_format.space_after = Pt(2)
font(author.add_run("Project Team"), 11)
aff = doc.add_paragraph()
aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
aff.paragraph_format.space_after = Pt(8)
font(aff.add_run("Department of Computer Science and Engineering\nIndia"), 10, italic=True)

abstract = doc.add_paragraph()
abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
abstract.paragraph_format.left_indent = Inches(0.45)
abstract.paragraph_format.right_indent = Inches(0.45)
abstract.paragraph_format.space_after = Pt(4)
font(abstract.add_run("Abstract—"), 9, bold=True, italic=True)
font(abstract.add_run(
    "The increasing accessibility of generative media tools has intensified the need for practical systems that screen both manipulated facial imagery and misleading textual content. This paper presents a lightweight multimodal detection framework comprising an EfficientNet-based binary image classifier, a term-frequency–inverse-document-frequency (TF-IDF) and logistic-regression news classifier, a Flask application programming interface, and Android and web clients. The image pipeline uses stratified train, validation, and test partitions, transfer learning, augmentation, fine-tuning, threshold calibration, and uncertainty rejection. The completed V2 image model achieved 80.41% balanced accuracy on a held-out test set, with fake and real recalls of 76.09% and 84.73%, respectively. The text model was evaluated by holding out an entire source; it achieved 65.4% accuracy but only 0.481 macro-F1, demonstrating substantial cross-source distribution shift. These results show that a deployable screening tool can be constructed using moderate computing resources, while also exposing the limits of closed-set classification. The system therefore reports uncertain predictions for low-confidence inputs and explicitly positions outputs as decision support rather than forensic proof."), 9, italic=True)

kw = doc.add_paragraph()
kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
kw.paragraph_format.left_indent = Inches(0.45)
kw.paragraph_format.right_indent = Inches(0.45)
kw.paragraph_format.space_after = Pt(4)
font(kw.add_run("Index Terms—"), 9, bold=True, italic=True)
font(kw.add_run("deepfake detection, fake news, EfficientNet, TF-IDF, transfer learning, uncertainty, Android."), 9, italic=True)

section = doc.add_section(WD_SECTION.CONTINUOUS)
configure_section(section, 2)

add_heading(doc, "I. Introduction")
add_body(doc, "Synthetic facial imagery and rapidly propagated misinformation threaten trust in digital communication. Modern manipulation systems can generate visually plausible faces, while misleading news may imitate the lexical style of legitimate reporting. The two problems are related at the application level but require different feature representations: spatial visual artifacts for images and statistical linguistic patterns for text. FaceForensics++ demonstrated both the scale of facial manipulation research and the sensitivity of detectors to manipulation method and compression [1]. FakeNewsNet similarly emphasized that news content alone is insufficient for comprehensive misinformation analysis because social context and propagation behavior also matter [2].")
add_body(doc, "This work develops a practical screening application that accepts either a face image or news text and returns a calibrated result through web and Android interfaces. The contribution is not a claim of universal forensic detection. Instead, the work integrates independently trained visual and textual models behind a common service, evaluates them with class-balanced and cross-source protocols, and introduces an explicit manual-review state for weak predictions. The primary contributions are: (1) a resource-conscious image pipeline using transfer learning and calibrated binary decisions; (2) a source-aware text evaluation that exposes dataset shift; (3) a Flask REST interface shared by web and Android clients; and (4) a V3 hard-example-mining design intended to improve difficult visual cases without falsifying labels or hard-coding samples.")

add_heading(doc, "II. Related Work")
add_subheading(doc, "A. Visual Manipulation Detection")
add_body(doc, "Rossler et al. introduced FaceForensics++, a benchmark containing pristine content and multiple manipulation families, and showed that domain knowledge and face-focused processing materially improve detection [1]. MobileNetV2 uses inverted residuals and linear bottlenecks to support efficient inference on constrained devices [3]. EfficientNet scales network depth, width, and resolution jointly [4], while EfficientNetV2 further optimizes training speed and parameter efficiency through training-aware architecture search and fused mobile inverted bottlenecks [5]. These architectures motivate the progression from MobileNetV2 and EfficientNet-B0 baselines to the EfficientNetV2-B0 design used in the pending V3 experiment.")
add_subheading(doc, "B. Misinformation and Fact Verification")
add_body(doc, "FakeNewsNet combines article content with social and spatiotemporal context [2]. FEVER contains 185,445 human-generated claims labeled as supported, refuted, or not enough information and associates evidence with supported and refuted claims [6]. Fakeddit contributes more than one million multimodal samples and fine-grained labels [7]. The present text classifier intentionally remains a pattern detector rather than an evidence-retrieval fact checker. This distinction is essential: TF-IDF and logistic regression can learn source-correlated language but cannot establish whether a current factual claim is true.")

doc.add_page_break()
add_heading(doc, "III. System Architecture")
make_figure()
make_flowcharts()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
picture = p.add_run().add_picture(str(FIG), width=Inches(3.25))
doc_pr = picture._inline.docPr
doc_pr.set("descr", "Architecture diagram showing separate image and news pipelines feeding a Flask API and Android or web clients.")
add_caption(doc, "Fig. 1. End-to-end architecture of the dual-input screening system.")
add_body(doc, "Figure 1 shows two independent inference paths. Image files are decoded to RGB, resized, and evaluated by the convolutional model. Text is normalized and transformed into sparse TF-IDF vectors before classification. Both outputs are exposed through Flask endpoints. The Android client uploads images as multipart form data and news as JSON. The server also produces a visual artifact map; this map highlights local high-frequency inconsistencies but is not a causal explanation and must not be interpreted as proof of manipulation.")

add_heading(doc, "IV. Methodology")
add_subheading(doc, "A. Image Dataset and Partitioning")
add_body(doc, "The local image corpus contains 22,044 face images: 10,961 labeled fake and 11,083 labeled real. Paths are divided using two stratified splits with a fixed seed of 42, resulting in approximately 70% training, 15% validation, and 15% testing data. Stratification prevents class-order artifacts and preserves the nearly balanced label distribution. Images in V2 are resized to 160 x 160 pixels. Augmentation includes horizontal flipping, small rotations, zoom, and contrast variation.")
add_project_picture(doc, TRAINING_FLOW_FIG, "Fig. 2. Image-model training and evaluation workflow.",
                    "Flowchart from labeled data collection through stratified splitting, preprocessing, training, calibration, and testing.")

doc.add_page_break()
add_subheading(doc, "B. Transfer Learning and Calibration")
add_body(doc, "V2 uses an ImageNet-pretrained EfficientNet-B0 backbone. The backbone is initially frozen while a global-average-pooling layer, dropout, and sigmoid output are trained. The last 25 backbone layers are subsequently fine-tuned at a lower learning rate. Binary cross-entropy is optimized, with validation area under the receiver operating characteristic curve controlling checkpoint selection and early stopping. The model emits p_r, interpreted as the probability of the real class. A decision threshold tau is selected on the validation set by maximizing balanced accuracy:")
add_equation(doc, "BA(τ) = 1/2 [Recall_fake(τ) + Recall_real(τ)]", 1)
add_body(doc, "Because a calibrated threshold does not guarantee that an individual probability is reliable, predictions are rejected when the maximum class probability is below 0.65 or when the score lies near the selected threshold. This policy converts weak decisions into a manual-review outcome instead of presenting them as definitive.")
add_subheading(doc, "C. V3 Hard-Example Design")
add_body(doc, "The V3 pipeline raises resolution to 224 x 224 pixels, preserves aspect ratio with padding, adopts EfficientNetV2-B0, and uses binary focal loss. Focal loss reduces the contribution of easy examples and emphasizes difficult ones [8]. After conventional training and fine-tuning, the training set is rescored; misclassified samples receive four times the normal sample weight during a final low-rate optimization stage. V3 results are deliberately omitted because training was still in progress when this manuscript was prepared.")
add_project_picture(doc, DECISION_FLOW_FIG, "Fig. 3. Runtime prediction and uncertainty-handling workflow.",
                    "Flowchart from user input through preprocessing and model scoring to a confident label or manual review.")

doc.add_page_break()
add_subheading(doc, "D. Text Classification")
add_body(doc, "The prepared news corpus contains 186,032 records from FakeNewsNet, FEVER-derived data, and WELFake. Text is normalized, filtered to at least 20 characters, and represented with lowercase word unigrams and bigrams. The TF-IDF vocabulary is capped at 100,000 features, with sublinear term frequency and English stop-word removal. A class-balanced logistic regression classifier estimates fake and real probabilities. For feature t in document d, the representation is:")
add_equation(doc, "w(t,d) = [1 + log tf(t,d)] log [N / df(t)]", 2)
add_body(doc, "Evaluation uses GroupShuffleSplit with the dataset source as the group, ensuring that all records from the held-out source are absent from training. After evaluation, a production artifact is fitted on all available records. Scores between 0.35 and 0.65 return a needs-fact-check result.")

add_heading(doc, "V. Experimental Results")
add_caption(doc, "TABLE I\nIMAGE MODEL RESULTS (%)")
add_table(doc, ["Model", "Split", "Bal. Acc.", "Fake Rec.", "Real Rec."], [
    ("Three-model ensemble", "Validation", "77.50", "70.30", "84.70"),
    ("EfficientNet-B0 V2", "Validation", "81.48", "—", "—"),
    ("EfficientNet-B0 V2", "Test", "80.41", "76.09", "84.73"),
], [1.05, 0.62, 0.62, 0.62, 0.62])
add_body(doc, "Table I reports stored metadata from completed experiments. V2 improves validation balanced accuracy over the earlier three-model ensemble, although comparisons should be interpreted cautiously because the pipelines use different split implementations. On the V2 test set, real recall exceeds fake recall by 8.64 percentage points, indicating that false negatives remain the dominant class-specific weakness.")
add_caption(doc, "TABLE II\nCROSS-SOURCE TEXT RESULTS")
add_table(doc, ["Metric", "Fake", "Real", "Overall"], [
    ("Precision", "0.215", "0.748", "—"),
    ("Recall", "0.155", "0.817", "—"),
    ("F1-score", "0.180", "0.781", "0.481 macro"),
    ("Accuracy", "—", "—", "0.654"),
], [1.05, 0.68, 0.68, 0.99])
add_body(doc, "FakeNewsNet was held out during source-aware evaluation. Table II shows that overall accuracy is inflated by the larger real class, whereas macro-F1 exposes weak fake-news transfer. The fake recall of 0.155 means that the lexical model misses most fake items from the unseen source. This result supports uncertainty rejection and confirms that the module should be described as a text-pattern screener, not a truth engine.")

doc.add_page_break()
add_heading(doc, "VI. Deployment")
add_body(doc, "The Flask service provides POST /api/predict-image and POST /api/predict-news endpoints. The image endpoint stores a randomized upload name, runs inference, creates an artifact visualization, and returns JSON. The news endpoint validates JSON text before classification. The Android application uses the system image picker, previews the selected image, sends requests on a background thread, displays progress states, and stores the chosen server address. During emulator testing, 10.0.2.2 maps to the host machine. A physical device must use the host's local network address.")
add_body(doc, "The design separates model artifacts from interface code, allowing V3 to replace V2 automatically when both its model and metadata files exist. Cleartext HTTP is enabled for local development only; production deployment should use HTTPS, authentication, upload-size limits, retention controls, and rate limiting.")

add_heading(doc, "VII. Limitations and Ethical Considerations")
add_body(doc, "The system has four important limitations. First, benchmark accuracy does not guarantee performance on new generators, recompressed social-media images, partial faces, or adversarial transformations. Second, a dataset label may itself be noisy. Third, a content-only news model cannot retrieve evidence or reason about current events. Fourth, the visualization is based on local detail energy rather than a validated causal attribution method. Consequently, outputs must not be used as sole evidence for disciplinary, legal, journalistic, or financial decisions.")
add_body(doc, "The observed false negative on a known fake training image illustrates why hard-coded corrections are inappropriate. Mapping a filename or hash to its stored label would improve a demonstration but not the detector. The chosen response is to report low-confidence cases as uncertain and to improve the learning procedure through higher resolution, focal loss, hard-example weighting, and future cross-dataset evaluation.")

doc.add_page_break()
add_heading(doc, "VIII. Conclusion and Future Work")
add_body(doc, "This paper presented a lightweight end-to-end framework for screening manipulated face images and suspicious news text. The completed image model reached 80.41% test balanced accuracy, while cross-source text performance remained limited at 0.481 macro-F1. The contrast between in-domain image results and source-shifted text results demonstrates that honest evaluation and uncertainty communication are as important as model selection. Future work will complete and benchmark V3, add face detection and frequency-domain features, evaluate on external deepfake datasets, incorporate evidence retrieval for claims, calibrate probabilities with reliability diagrams, and migrate the service to authenticated HTTPS deployment.")

add_heading(doc, "References")
refs = [
    "[1] A. Rossler, D. Cozzolino, L. Verdoliva, C. Riess, J. Thies, and M. Niessner, “FaceForensics++: Learning to detect manipulated facial images,” in Proc. IEEE/CVF Int. Conf. Comput. Vis. (ICCV), 2019, pp. 1–11.",
    "[2] K. Shu, D. Mahudeswaran, S. Wang, D. Lee, and H. Liu, “FakeNewsNet: A data repository with news content, social context and spatiotemporal information for studying fake news on social media,” arXiv:1809.01286, 2018.",
    "[3] M. Sandler, A. Howard, M. Zhu, A. Zhmoginov, and L.-C. Chen, “MobileNetV2: Inverted residuals and linear bottlenecks,” in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), 2018, pp. 4510–4520.",
    "[4] M. Tan and Q. V. Le, “EfficientNet: Rethinking model scaling for convolutional neural networks,” in Proc. 36th Int. Conf. Mach. Learn. (ICML), vol. 97, 2019, pp. 6105–6114.",
    "[5] M. Tan and Q. Le, “EfficientNetV2: Smaller models and faster training,” in Proc. 38th Int. Conf. Mach. Learn. (ICML), vol. 139, 2021, pp. 10096–10106.",
    "[6] J. Thorne, A. Vlachos, C. Christodoulopoulos, and A. Mittal, “FEVER: A large-scale dataset for fact extraction and verification,” in Proc. NAACL-HLT, 2018, pp. 809–819.",
    "[7] K. Nakamura, S. Levy, and W. Y. Wang, “r/Fakeddit: A new multimodal benchmark dataset for fine-grained fake news detection,” in Proc. 12th Lang. Resour. Eval. Conf. (LREC), 2020, pp. 6149–6157.",
    "[8] T.-Y. Lin, P. Goyal, R. Girshick, K. He, and P. Dollar, “Focal loss for dense object detection,” in Proc. IEEE Int. Conf. Comput. Vis. (ICCV), 2017, pp. 2980–2988.",
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(ref), 8)

props = doc.core_properties
props.title = "A Lightweight Multimodal Framework for Deepfake Image and Fake-News Detection"
props.subject = "IEEE-style research paper"
props.author = "Project Team"
props.keywords = "deepfake detection; fake news; EfficientNet; TF-IDF; Android"
doc.save(OUT)
print(OUT)
